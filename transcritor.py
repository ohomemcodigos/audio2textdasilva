import os
import time
import threading
import sys
import re
from faster_whisper import WhisperModel
from deep_translator import GoogleTranslator
import docx

class InterceptadorDownload:
    def __init__(self, on_status):
        self.on_status = on_status
        self.original_stderr = sys.stderr

    def write(self, msg):
        match = re.search(r'(\d+)%', msg)
        if match:
            self.on_status(f"A descarregar o modelo de IA... {match.group(1)}%")
        self.original_stderr.write(msg)

    def flush(self):
        self.original_stderr.flush()

# formatador de tempo para o padrao srt (HH:MM:SS,mmm)
def formatar_tempo_srt(segundos):
    milisegundos = int((segundos - int(segundos)) * 1000)
    s = int(segundos)
    m, s = divmod(s, 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d},{milisegundos:03d}"

# formatador de tempo padrao (HH:MM:SS)
def formatar_tempo_txt(segundos):
    return time.strftime('%H:%M:%S', time.gmtime(segundos))

def processar_audios_local(lista_destinos, modelo, idioma, traduzir, formato, estilo, threads, on_status, on_progress, on_finish, on_error):
    def tarefa():
        interceptador = InterceptadorDownload(on_status)
        try:
            on_status(f"A preparar o carregamento do modelo {modelo}...")
            sys.stderr = interceptador
            
            # aplica o fine-tuning de CPU definido nas configuracoes
            model = WhisperModel(modelo, device="cpu", compute_type="int8", cpu_threads=threads)
            sys.stderr = interceptador.original_stderr
            
            mapa_idiomas = {"Inglês": "en", "Português": "pt", "Espanhol": "es"}
            lang_code = mapa_idiomas.get(idioma, None) 
            
            total = len(lista_destinos)
            
            for index, (caminho_audio, caminho_destino_base) in enumerate(lista_destinos):
                nome_base = os.path.basename(caminho_audio)
                on_status(f"Transcrevendo {index+1}/{total}: {nome_base} (A iniciar...)")
                
                segments, info = model.transcribe(caminho_audio, language=lang_code, beam_size=5)
                duracao_total = info.duration
                
                texto_linhas = []
                inicio_tempo = time.time()
                
                for i, segment in enumerate(segments):
                    texto_traduzido = segment.text.strip()
                    if traduzir:
                        texto_traduzido = GoogleTranslator(source='auto', target='pt').translate(texto_traduzido)
                    
                    # formata de acordo com a escolha do utilizador
                    if formato == ".srt":
                        tempo_inicio = formatar_tempo_srt(segment.start)
                        tempo_fim = formatar_tempo_srt(segment.end)
                        bloco_srt = f"{i + 1}\n{tempo_inicio} --> {tempo_fim}\n{texto_traduzido}\n"
                        texto_linhas.append(bloco_srt)
                    else:
                        if estilo == "com_tempo":
                            tempo = formatar_tempo_txt(segment.start)
                            texto_linhas.append(f"[{tempo}] {texto_traduzido}")
                        else:
                            # texto corrido
                            texto_linhas.append(f"{texto_traduzido} ")
                    
                    tempo_decorrido = time.time() - inicio_tempo
                    progresso_ficheiro = segment.end / duracao_total if duracao_total > 0 else 0
                    
                    if progresso_ficheiro > 0 and progresso_ficheiro <= 1.0:
                        tempo_total_estimado = tempo_decorrido / progresso_ficheiro
                        tempo_restante = max(0, tempo_total_estimado - tempo_decorrido)
                        
                        eta_str = time.strftime('%H:%M:%S' if tempo_restante > 3600 else '%M:%S', time.gmtime(tempo_restante))
                        percentagem = int(progresso_ficheiro * 100)
                        on_status(f"Transcrevendo {index+1}/{total}: {nome_base} | {percentagem}% | Restante: {eta_str}")
                    
                    progresso_global = (index / total) + (progresso_ficheiro / total)
                    on_progress(progresso_global)
                
                # criacao final do ficheiro na extensao correta
                caminho_final = f"{caminho_destino_base}{formato}"
                
                if formato == ".docx":
                    doc = docx.Document()
                    paragrafo = doc.add_paragraph()
                    if estilo == "corrido":
                        paragrafo.add_run("".join(texto_linhas))
                    else:
                        for linha in texto_linhas:
                            doc.add_paragraph(linha)
                    doc.save(caminho_final)
                else:
                    with open(caminho_final, "w", encoding="utf-8") as f:
                        if formato == ".srt" or estilo == "com_tempo":
                            f.write("\n".join(texto_linhas))
                        else:
                            f.write("".join(texto_linhas))
                    
            on_finish()
            
        except Exception as e:
            if hasattr(interceptador, 'original_stderr'):
                sys.stderr = interceptador.original_stderr
            on_error(str(e))

    threading.Thread(target=tarefa, daemon=True).start()